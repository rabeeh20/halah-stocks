"""
Export All NSE Listed Stocks to Excel
======================================
Run this script on your Mac to download the complete list of
all NSE-listed equity stocks and save it as an Excel file.

Usage:
    cd /Users/rabeehvailassery/Desktop/halal\ stocks/scripts
    pip3 install requests pandas openpyxl
    python3 export_all_nse_stocks.py

Output:
    ~/Desktop/NSE_All_Stocks.xlsx
"""

import requests
import pandas as pd
import io
import os
from datetime import datetime

# ── Output file path ──────────────────────────────────────────────────────────
OUTPUT_PATH = os.path.expanduser("~/Desktop/NSE_All_Stocks.xlsx")

# ── NSE provides EQUITY_L.csv — the full list of all listed equities ──────────
NSE_EQUITY_LIST_URL = "https://archives.nseindia.com/content/equities/EQUITY_L.csv"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.9",
    "Referer": "https://www.nseindia.com/",
}


def fetch_nse_equity_list() -> pd.DataFrame:
    """Download the complete NSE equity list CSV and return as DataFrame."""
    print("📡 Fetching NSE equity list from official NSE archives...")
    print(f"   URL: {NSE_EQUITY_LIST_URL}\n")

    session = requests.Session()
    session.headers.update(HEADERS)

    # Warm up the session with NSE homepage first
    try:
        print("   Warming up NSE session...")
        session.get("https://www.nseindia.com", timeout=15)
    except Exception:
        pass

    # Now fetch the CSV
    response = session.get(NSE_EQUITY_LIST_URL, timeout=30)
    response.raise_for_status()

    df = pd.read_csv(io.StringIO(response.text))

    # Clean up column names (remove extra spaces)
    df.columns = [col.strip() for col in df.columns]

    # Strip whitespace from all string columns
    for col in df.select_dtypes(include="object").columns:
        df[col] = df[col].str.strip()

    print(f"✅ Fetched {len(df):,} stocks from NSE.\n")
    return df


def export_to_excel(df: pd.DataFrame, output_path: str):
    """Export the DataFrame to a nicely formatted Excel file."""
    print(f"💾 Saving to Excel: {output_path}")

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="NSE All Stocks")

        # ── Auto-size columns ──────────────────────────────────────────────
        worksheet = writer.sheets["NSE All Stocks"]

        for col_idx, col in enumerate(df.columns, start=1):
            max_len = max(
                len(str(col)),
                df[col].astype(str).str.len().max() if not df[col].empty else 0
            )
            # Cap at 50 characters wide
            col_letter = worksheet.cell(row=1, column=col_idx).column_letter
            worksheet.column_dimensions[col_letter].width = min(max_len + 4, 50)

        # ── Freeze top header row ──────────────────────────────────────────
        worksheet.freeze_panes = "A2"

        # ── Bold the header row ────────────────────────────────────────────
        from openpyxl.styles import Font, PatternFill, Alignment
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center")

        for cell in worksheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            worksheet.row_dimensions[1].height = 22

    print(f"✅ Excel file saved successfully!\n")


def main():
    print("=" * 60)
    print("  NSE Complete Stock List Exporter")
    print(f"  Date: {datetime.now().strftime('%d %B %Y, %I:%M %p')}")
    print("=" * 60)
    print()

    try:
        df = fetch_nse_equity_list()

        # Show a summary of what we got
        print("📊 Column Overview:")
        for col in df.columns:
            print(f"   • {col}")
        print()

        print(f"📈 Total Stocks Listed: {len(df):,}")

        if "SERIES" in df.columns:
            series_counts = df["SERIES"].value_counts()
            print("\n📂 Breakdown by Series:")
            for series, count in series_counts.items():
                print(f"   {series}: {count:,} stocks")

        print()

        export_to_excel(df, OUTPUT_PATH)

        print("=" * 60)
        print(f"✅ Done! File saved at:")
        print(f"   {OUTPUT_PATH}")
        print("=" * 60)

    except requests.exceptions.HTTPError as e:
        print(f"❌ HTTP Error: {e}")
        print("   NSE may be blocking the request. Try running again after a few minutes.")
    except Exception as e:
        print(f"❌ Error: {e}")
        raise


if __name__ == "__main__":
    main()
